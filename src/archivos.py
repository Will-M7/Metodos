"""Extracción con procedencia y sin modificar originales."""
import hashlib
import json
from pathlib import Path
import zipfile

SUPPORTED = {'.pdf', '.xlsx', '.xlsm', '.m', '.txt', '.md', '.csv', '.png', '.jpg', '.jpeg'}
IMAGES = {'.png', '.jpg', '.jpeg'}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def discover(folder):
    if not folder.exists():
        return []
    return sorted(p for p in folder.rglob('*') if p.is_file() and not p.is_symlink()
                  and not any(x.startswith('.') for x in p.relative_to(folder).parts)
                  and not any(x.is_symlink() for x in p.parents if x != folder.parent))


def inputs(root, mode):
    paths = discover(root / 'entrada')
    paths += [p for p in root.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED
              and p.name not in {'README.md', 'AGENTS.md', 'requirements.txt'} and not p.is_symlink()]
    if mode == 'revisar':
        paths += discover(root / 'resultados')
    return sorted(set(paths))


def extract(path, root, scratch):
    suffix = path.suffix.lower()
    result = {'archivo': path.relative_to(root).as_posix(), 'sha256': digest(path),
              'contenido': '', 'advertencias': [], 'imagenes': [], 'lectura_completa': True}
    if suffix in IMAGES:
        result['imagenes'] = [str(path.resolve())]
        result['contenido'] = 'Referencia visual adjunta. No deducir texto ilegible.'
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        texts = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            texts.append(f'Página {i}:\n{text}')
        result['contenido'] = '\n\n'.join(texts)
        # Renderizar también preserva fórmulas, tablas y páginas escaneadas.
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(str(path))
            try:
                for i in range(len(pdf)):
                    target = scratch / f'{result["sha256"]}_{i+1}.png'
                    page = pdf[i]
                    bitmap = page.render(scale=1.5)
                    bitmap.to_pil().save(target)
                    result['imagenes'].append(str(target))
                    bitmap.close()
                    page.close()
            finally:
                pdf.close()
        except ImportError:
            result['advertencias'].append('PDF sin inspección visual: instalar pypdfium2 y Pillow; fórmulas/tablas pueden perder estructura.')
    elif suffix in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False)
        cached = load_workbook(path, data_only=True, read_only=True)
        try:
            sheets = []
            for ws in wb:
                if ws.max_row * ws.max_column > 150000:
                    raise ValueError(f'Hoja demasiado grande para este MVP: {ws.title}')
                cells = []
                for row in ws:
                    for cell in row:
                        if cell.value is not None:
                            value = cell.value
                            if hasattr(value, 'text'):
                                value = {'formula_matricial': value.text, 'rango': value.ref}
                            cells.append({'celda': cell.coordinate, 'valor': str(value),
                                          'cache_no_verificada': str(cached[ws.title][cell.coordinate].value),
                                          'formato': cell.number_format,
                                          'negrita': bool(cell.font.bold),
                                          'alineacion': cell.alignment.horizontal})
                sheets.append({'hoja': ws.title, 'celdas': cells,
                               'combinadas': [str(r) for r in ws.merged_cells.ranges],
                               'anchos': {k: d.width for k, d in ws.column_dimensions.items()}})
            result['contenido'] = json.dumps(sheets, ensure_ascii=False)
        finally:
            wb.close()
            cached.close()
        with zipfile.ZipFile(path) as z:
            for i, name in enumerate(z.namelist()):
                if name.startswith('xl/media/') and Path(name).suffix.lower() in IMAGES:
                    target = scratch / f'{result["sha256"]}_{i}{Path(name).suffix}'
                    target.write_bytes(z.read(name))
                    result['imagenes'].append(str(target))
    elif suffix in SUPPORTED:
        try:
            result['contenido'] = path.read_text(encoding='utf-8-sig')
        except UnicodeDecodeError:
            result['contenido'] = path.read_text(encoding='cp1252')
            result['advertencias'].append('Texto leído como Windows-1252; confirmar símbolos.')
    else:
        result['lectura_completa'] = False
        result['advertencias'].append('Formato no compatible; exportar a PDF, XLSX, M o imagen.')
    return result


def pack(records):
    text = json.dumps([{k: v for k, v in r.items() if k != 'imagenes'} for r in records], ensure_ascii=False)
    images = list(dict.fromkeys(i for r in records for i in r['imagenes']))
    if len(text) > 220000 or len(images) > 50:
        raise ValueError('Material excede el contexto del MVP (220000 caracteres / 50 imágenes). Dividir el trabajo; no se omitió contenido silenciosamente.')
    return text, images


def safe_extract(path, root, scratch):
    try:
        return extract(path, root, scratch)
    except Exception as e:
        return {'archivo': path.relative_to(root).as_posix(), 'sha256': digest(path),
                'contenido': '', 'advertencias': ['NO LEÍDO: ' + str(e)],
                'imagenes': [], 'lectura_completa': False}
