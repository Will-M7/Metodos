"""Libro académico monocromo; fórmulas comprobadas y valores cacheados reales."""
import re
import xml.etree.ElementTree as ET
import zipfile
from .verificacion import FormulaBook, close


def build_sheet(wb, q, template=False):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    name = q['hoja']
    if not name or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name):
        raise ValueError('Nombre de hoja inválido')
    if template:
        if q['plantilla_hoja'] not in wb.sheetnames:
            raise ValueError('Falta hoja de plantilla: ' + q['plantilla_hoja'])
        ws = wb[q['plantilla_hoja']]
        # No renombrar plantillas: otras fórmulas podrían depender de su nombre.
    else:
        if name in wb.sheetnames:
            raise ValueError('Dos ejercicios no pueden compartir hoja')
        ws = wb.create_sheet(name)
    assigned = set()
    for item in q['celdas']:
        addr = item['celda'].upper()
        if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,4}', addr) or addr in assigned:
            raise ValueError('Celda inválida/duplicada: ' + addr)
        assigned.add(addr)
        c = ws[addr]
        from openpyxl.cell.cell import MergedCell
        if isinstance(c, MergedCell):
            raise ValueError('Escribir en la esquina superior izquierda de la celda combinada: ' + addr)
        if c.row > 20000 or c.column > 100:
            raise ValueError('Hoja fuera de límites del MVP')
        value = item['valor']
        if item['tipo'] == 'numero':
            value = float(value)
        elif item['tipo'] == 'formula':
            if not value.startswith('='):
                raise ValueError('Fórmula sin =')
        elif value.startswith(('=', '+', '-', '@')):
            # Forzar texto literal, no convertir una nota en fórmula.
            c.value = value
            c.data_type = 's'
        c.value = value
        if item['tipo'] == 'texto':
            c.data_type = 's'
        if not template:
            c.font = Font(name='Calibri', size=11, color='000000', bold=c.row == 1)
            c.fill = PatternFill(fill_type='solid', fgColor='FFFFFF')
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if item['tipo'] != 'texto':
                c.number_format = '0.0000000000'
            ws.column_dimensions[c.column_letter].width = max(18, min(65, len(str(value)) + 2))
            ws.row_dimensions[c.row].height = max(ws.row_dimensions[c.row].height or 15, 30 if len(str(value)) > 45 else 16)
    evaluator = FormulaBook(ws)
    formula_count = 0
    for row in ws:
        for cell in row:
            if cell.data_type == 'f':
                evaluator.value(cell.coordinate)
                formula_count += 1
    if not formula_count:
        raise ValueError('El desarrollo no contiene fórmulas Excel')
    for item in q['celdas']:
        if item['tipo'] in {'formula', 'numero'}:
            if item['esperado'] is None or not close(evaluator.value(item['celda']), item['esperado'], 1e-7):
                raise ValueError('Valor de celda no coincide: ' + item['celda'])
    for out in q['salidas']:
        if len(out['celdas']) != len(out['valores']) or not out['celdas']:
            raise ValueError('Cada salida Octave necesita celdas Excel correspondientes')
        for addr, val in zip(out['celdas'], out['valores']):
            if not close(evaluator.value(addr), val, out['tolerancia']):
                raise ValueError('Excel difiere de salida Octave prevista: ' + addr)
    return ws, evaluator.cache


def save(wb, path, caches):
    """openpyxl no calcula. Insertar sólo valores que nuestro evaluador comprobó."""
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    wb.save(path)
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    temp = path.with_suffix('.tmp')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(temp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            match = re.fullmatch(r'xl/worksheets/sheet([0-9]+).xml', item.filename)
            if match:
                index = int(match.group(1)) - 1
                cache = caches.get(wb.worksheets[index].title, {})
                tree = ET.fromstring(data)
                for cell in tree.findall('.//s:c', ns):
                    addr = cell.attrib['r']
                    if cell.find('s:f', ns) is not None and addr in cache:
                        val = cell.find('s:v', ns)
                        if val is None:
                            val = ET.SubElement(cell, '{' + ns['s'] + '}v')
                        val.text = repr(cache[addr])
                data = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            zout.writestr(item, data)
    temp.replace(path)
    from openpyxl import load_workbook
    check = load_workbook(path, data_only=True)
    try:
        for name, cells in caches.items():
            for addr, value in cells.items():
                actual = check[name][addr].value
                if actual is None and wb[name][addr].value is None:
                    actual = 0
                if not close(actual, value, 1e-7):
                    raise ValueError('Falló relectura del Excel: ' + name + '!' + addr)
    finally:
        check.close()


def load_template(path):
    from openpyxl import load_workbook
    # openpyxl pierde algunos objetos. Rechazar plantillas no preservables, sin tocar el original.
    with zipfile.ZipFile(path) as z:
        unsupported = [n for n in z.namelist() if n.startswith(('xl/drawings/', 'xl/activeX/', 'xl/ctrlProps/', 'xl/slicer'))
                       or n.endswith('vbaProject.bin')]
        if unsupported:
            raise ValueError('Plantilla contiene objetos que openpyxl no preserva con garantía; pendiente adaptar sin pérdida: ' + path.name)
    return load_workbook(path)
