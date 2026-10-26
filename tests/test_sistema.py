"""Pruebas reales locales y dobles explícitos, siempre fuera de los datos del usuario."""
import contextlib
from copy import deepcopy
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from src import archivos, codex, excel, main, memoria, octave
from src.verificacion import arithmetic, check_question, FormulaBook

ROOT = Path(__file__).resolve().parents[1]


def question():
    return {
        'numero': '1', 'metodo': 'Sistema lineal de prueba', 'base': 'explicita',
        'fuentes': ['entrada/problema.txt'], 'estado': 'resuelta', 'respuesta': 'x1 = 1; x2 = 1',
        'hoja': 'Pregunta 1', 'plantilla_hoja': '',
        'celdas': [
            {'celda': 'A1', 'tipo': 'texto', 'valor': 'Sistema lineal', 'esperado': None},
            {'celda': 'A2', 'tipo': 'texto', 'valor': 'x1', 'esperado': None},
            {'celda': 'B2', 'tipo': 'numero', 'valor': '1', 'esperado': 1},
            {'celda': 'A3', 'tipo': 'texto', 'valor': 'x2 = (6-4*x1)/2', 'esperado': None},
            {'celda': 'B3', 'tipo': 'formula', 'valor': '=(6-4*B2)/2', 'esperado': 1},
            {'celda': 'A4', 'tipo': 'texto', 'valor': 'Residuo segunda ecuación', 'esperado': None},
            {'celda': 'B4', 'tipo': 'formula', 'valor': '=2*B2+3*B3-5', 'esperado': 0}
        ],
        'octave': 'clc\nclear\nA = [4 2; 2 3]\nb = [6; 5]\nx = A\\b\nresiduo = A*x-b\n',
        'comprobaciones': [
            {'nombre': 'Primera ecuación original', 'expresion': '4*1+2*1-6', 'esperado': 0, 'tolerancia': 1e-8},
            {'nombre': 'Segunda ecuación original', 'expresion': '2*1+3*1-5', 'esperado': 0, 'tolerancia': 1e-8}
        ],
        'salidas': [{'variable': 'x', 'valores': [1, 1], 'celdas': ['B2', 'B3'], 'tolerancia': 1e-8}],
        'requiere_grafica': False, 'requiere_captura': True, 'revision': 'CORRECTA', 'errores': []
    }


def work(q=None):
    return {'preguntas': [q or question()], 'completa': True, 'nota': 20,
            'nota_estimada': True, 'veredicto': 'Prueba con doble de Codex; estimación uniforme.', 'pendientes': []}


def fixture(root):
    for name in ('entrada', 'resultados', 'referencias_clase', 'perfil_academico'):
        (root / name).mkdir()
    for p in (ROOT / 'perfil_academico').glob('*.md'):
        shutil.copyfile(p, root / 'perfil_academico' / p.name)
    shutil.copyfile(ROOT / 'AGENTS.md', root / 'AGENTS.md')
    (root / 'entrada' / 'problema.txt').write_text('Resolver 4*x1+2*x2=6; 2*x1+3*x2=5. Trabajo alumno: x1=x2=1.', encoding='utf-8')


def write_pdf(path):
    # Fixture PDF mínimo, sin reportlab como dependencia de producción/test.
    stream = b'BT /F1 12 Tf 40 100 Td (Pregunta 1: x^2 - 2 = 0) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
               b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 150] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    pdf = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for i, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += f'{i} 0 obj\n'.encode() + body + b'\nendobj\n'
    xref = len(pdf)
    pdf += f'xref\n0 {len(objects)+1}\n0000000000 65535 f \n'.encode()
    for offset in offsets[1:]:
        pdf += f'{offset:010d} 00000 n \n'.encode()
    pdf += f'trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    path.write_bytes(pdf)


class NumericTests(unittest.TestCase):
    def test_original_equations(self):
        check_question(question())
        self.assertAlmostEqual(arithmetic('sqrt(4)+sum([1,2])'), 5)
        bad = question()
        bad['comprobaciones'][0]['expresion'] = '4*2+2*1-6'
        with self.assertRaises(ValueError):
            check_question(bad)

    def test_no_generated_python_execution(self):
        for expression in ('__import__("os").system("id")', '(1).__class__', '[x for x in [1]]', '2**1000', '"a"*100', '1/0', 'sqrt([1])'):
            with self.assertRaises((ValueError, SyntaxError)):
                arithmetic(expression)

    def test_octave_packages_no_installation_or_file_operations(self):
        octave.vet('pkg load statistics\nx = 1')
        for code in ('pkg install statistics', 'load personal.mat', 'system("echo x")', 'delete("archivo")'):
            with self.assertRaises(ValueError):
                octave.vet(code)

    def test_strict_structured_output(self):
        codex.validate(work(), codex.WORK)
        bad = work()
        bad['nota'] = 21
        with self.assertRaises(ValueError):
            codex.validate(bad, codex.WORK)
        bad = work()
        bad['extra'] = 'no'
        with self.assertRaises(ValueError):
            codex.validate(bad, codex.WORK)


class ExcelExtractionTests(unittest.TestCase):
    def test_pdf_excel_octave_extraction_and_originals(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pdf, xlsx, script = root/'problema.pdf', root/'trabajo.xlsx', root/'pregunta.m'
            write_pdf(pdf)
            wb = Workbook()
            wb.active['A1'] = '=2+3'
            wb.save(xlsx)
            script.write_text('clc\nclear\nx = 2+3\n')
            before = {p.name: archivos.digest(p) for p in (pdf,xlsx,script)}
            self.assertIn('x^2', archivos.extract(pdf,root,root)['contenido'])
            self.assertIn('=2+3', archivos.extract(xlsx,root,root)['contenido'])
            self.assertIn('x = 2+3', archivos.extract(script,root,root)['contenido'])
            self.assertEqual(before, {p.name: archivos.digest(p) for p in (pdf,xlsx,script)})

    def test_array_formula_preserved_as_text(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.formula import ArrayFormula
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);path=root/'matrices.xlsx'
            wb=Workbook();wb.active['C1']=ArrayFormula('C1:C2','=MMULT(A1:B2,D1:D2)');wb.save(path)
            text=archivos.extract(path,root,root)['contenido']
            self.assertIn('MMULT(A1:B2,D1:D2)',text)
            self.assertNotIn('object at',text)

    def test_workbook_formulas_cache_and_recalculation(self):
        from openpyxl import Workbook, load_workbook
        wb = Workbook()
        wb.remove(wb.active)
        ws, cache = excel.build_sheet(wb, question())
        with tempfile.TemporaryDirectory() as d:
            path = Path(d)/'salida.xlsx'
            excel.save(wb,path,{ws.title:cache})
            formulas = load_workbook(path, data_only=False)
            values = load_workbook(path, data_only=True)
            self.assertEqual(formulas.active['B3'].value,'=(6-4*B2)/2')
            self.assertEqual(values.active['B3'].value,1)
            self.assertEqual(values.active['B4'].value,0)
            for row in formulas.active:
                for c in row:
                    if c.value is not None:
                        self.assertEqual(c.font.color.rgb, '00000000')
                        self.assertEqual(c.fill.fgColor.rgb, '00FFFFFF')
            formulas.active['B2'] = 0.5
            self.assertAlmostEqual(FormulaBook(formulas.active).value('B3'),2)
            formulas.close()
            values.close()

    def test_ranges_functions_circular_unknown(self):
        from openpyxl import Workbook
        ws=Workbook().active
        ws['A1'],ws['A2']=2,3
        ws['B1']='=SUM($A$1:$A$2)+SQRT($A$1^2)+LOG10(100)'
        self.assertAlmostEqual(FormulaBook(ws).value('B1'),9)
        ws['C1']='=D1'
        ws['D1']='=C1'
        with self.assertRaises(ValueError):
            FormulaBook(ws).value('C1')
        ws['E1']='=VLOOKUP(A1,A1:B2,2,0)'
        with self.assertRaises(ValueError):
            FormulaBook(ws).value('E1')

    def test_template_simple_preserved_and_complex_refused(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment
        import zipfile
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'plantilla.xlsx'
            wb=Workbook()
            wb.active.title='Pregunta 1'
            wb.active['H8']='espacio manual'
            wb.active.merge_cells('H8:J9')
            wb.active.column_dimensions['H'].width=33
            wb.save(path)
            original=archivos.digest(path)
            loaded=excel.load_template(path)
            q=question();q['plantilla_hoja']='Pregunta 1'
            excel.build_sheet(loaded,q,True)
            self.assertEqual(loaded.active['H8'].value,'espacio manual')
            self.assertEqual(loaded.active.column_dimensions['H'].width,33)
            self.assertEqual(archivos.digest(path),original)
            with zipfile.ZipFile(path,'a') as z:
                z.writestr('xl/drawings/drawing1.xml','<drawing/>')
            with self.assertRaises(ValueError):
                excel.load_template(path)


class MemoryTests(unittest.TestCase):
    def test_incremental_sources_update_delete_no_invention(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);fixture(root)
            ref=root/'referencias_clase'/'docente.md'
            ref.write_text('Ejemplo docente confirmado: f(x)=x^2-2. Bisección usa c=(a+b)/2.')
            response={'ejercicios':[{'id':'1','localizacion':'línea 1','enunciado':'Punto medio de 1 y 2',
                'procedimiento':'c=(a+b)/2','formulas':['c=(a+b)/2'],'origen':'docente','clasificacion':'VALIDADA','errores':[],
                'comprobaciones':[{'nombre':'punto medio','expresion':'(1+2)/2','esperado':1.5,'tolerancia':1e-8}],
                'octave':'c = (1+2)/2','salidas':[{'variable':'c','valores':[1.5],'celdas':[],'tolerancia':1e-8}]}],
                'reglas':[{'categoria':'metodos','regla':'Usar c=(a+b)/2.',
                     'fuente':'referencias_clase/docente.md','evidencia':'c=(a+b)/2','ejercicio':'1'}],'advertencias':[]}
            with patch('src.memoria.invoke',return_value=response) as ai:
                memoria.academic(root,root)
                memoria.academic(root,root)
                self.assertEqual(ai.call_count,1)
                text=(root/'perfil_academico/metodos.md').read_text()
                self.assertIn('Fuente: referencias_clase/docente.md',text)
                ref.write_text('Nueva referencia incompatible: pendiente.')
                ai.return_value={'ejercicios':[],'reglas':[],'advertencias':['conflicto']}
                memoria.academic(root,root)
                self.assertEqual(ai.call_count,2)
                self.assertNotIn('Usar c=(a+b)/2.',(root/'perfil_academico/metodos.md').read_text())
                ref.unlink()
                memoria.academic(root,root)
                self.assertEqual(ai.call_count,2)


@unittest.skipUnless(shutil.which('octave-cli') or shutil.which('octave'), 'GNU Octave no disponible')
class RealOctaveTests(unittest.TestCase):
    def test_real_octave_and_wrong_answer(self):
        q=question()
        self.assertIn(octave.MARK, main.verify(q))
        bad=deepcopy(q['salidas']);bad[0]['valores'][0]=2
        with self.assertRaises(ValueError):
            octave.run(q['octave'],bad)

    def test_graph_export_from_real_octave(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            code = "clc\nclear\nx = [0 1 2]\ny = x.^2\nplot(x,y)\nxlabel('x')\nylabel('y')\n"
            outputs = [{'variable':'y','valores':[0,1,4],'celdas':[],'tolerancia':1e-8}]
            octave.run(code, outputs, folder)
            from PIL import Image
            path = folder/'grafica_01.png'
            self.assertTrue(path.is_file())
            with Image.open(path) as image:
                self.assertGreater(image.width, 100)
                self.assertGreater(image.height, 100)

    def test_generation_final_qa_no_overwrite(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);fixture(root)
            q=question();q['requiere_captura']=False
            qa={'preguntas':[{'numero':'1','correcta':True,'captura_legible':False,'observaciones':[]}]}
            with patch('src.codex.invoke', return_value=qa), contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main.generate(root,work(q),[]),0)
                original=archivos.digest(root/'resultados/trabajo_resuelto.xlsx')
                self.assertEqual(main.generate(root,work(q),[]),0)
            self.assertEqual(original,archivos.digest(root/'resultados/trabajo_resuelto.xlsx'))
            self.assertEqual(len(list((root/'resultados').glob('ejecucion_*'))),1)

    def test_mathematical_method_bisection_real(self):
        code='''clc
clear
a = 1
b = 2
for k = 1:30
  c = (a+b)/2;
  if (a^2-2)*(c^2-2) < 0
    b = c;
  else
    a = c;
  endif
endfor
c
residuo = c^2-2
'''
        log=octave.run(code,[{'variable':'c','valores':[2**0.5],'celdas':[],'tolerancia':1e-8}])
        self.assertIn(octave.MARK,log)

    def test_full_generation_partial_failure_keeps_success(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);fixture(root)
            bad=question();bad['numero']='2';bad['hoja']='Pregunta 2'
            bad['comprobaciones'][0]['expresion']='1/0'
            data=work();data['preguntas'].append(bad)
            originals={p:archivos.digest(p) for p in (root/'entrada').iterdir()}
            with patch('src.codex.invoke', return_value={'preguntas':[{'numero':'1','correcta':True,'captura_legible':False,'observaciones':[]}]}), contextlib.redirect_stdout(io.StringIO()) as terminal:
                code=main.generate(root,data,[])
            self.assertEqual(code,2)
            self.assertIn('P2: PENDIENTE',terminal.getvalue())
            path=root/'resultados/trabajo_resuelto.xlsx'
            self.assertTrue(path.exists())
            self.assertTrue((root/'resultados/octave/pregunta_01.salida.txt').exists())
            self.assertEqual(list((root/'resultados/capturas').glob('*.png')), [])
            wb=load_workbook(path,data_only=True)
            self.assertEqual(wb.sheetnames,['Pregunta 1'])
            self.assertEqual(wb.active['B3'].value,1)
            wb.close()
            self.assertEqual(originals,{p:archivos.digest(p) for p in (root/'entrada').iterdir()})

    def test_answers_and_review_end_to_end_with_explicit_ai_double(self):
        for mode in ('respuestas','revisar'):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as d:
                root=Path(d);fixture(root)
                with patch('src.codex.invoke',return_value=work()) as ai, contextlib.redirect_stdout(io.StringIO()) as terminal:
                    self.assertEqual(main.run(mode,root),0)
                self.assertEqual(ai.call_count,1)
                self.assertEqual(list((root/'resultados').iterdir()),[])
                self.assertIn('x1 = 1' if mode=='respuestas' else '20.00/20',terminal.getvalue())
                self.assertIn('MODO: '+mode,ai.call_args.args[1])

    def test_unreadable_input_retains_results_without_definitive_grade(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);fixture(root)
            (root/'entrada/incompleto.xlsx').write_bytes(b'archivo corrupto')
            with patch('src.codex.invoke',return_value=work()), contextlib.redirect_stdout(io.StringIO()) as terminal:
                self.assertEqual(main.run('revisar',root),2)
            self.assertNotIn('NOTA FINAL:',terminal.getvalue())
            self.assertIn('Archivos sin lectura completa',terminal.getvalue())

    def test_method_conflict_and_incomplete_review(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);fixture(root)
            data=work();data['preguntas'][0]['base']='conflicto'
            with patch('src.codex.invoke',return_value=data), contextlib.redirect_stdout(io.StringIO()) as terminal:
                code=main.run('revisar',root)
            self.assertEqual(code,2)
            self.assertIn('sin nota definitiva',terminal.getvalue())
            self.assertNotIn('NOTA FINAL:',terminal.getvalue())


class CodexBoundaryTests(unittest.TestCase):
    def test_large_course_selects_references_without_silent_truncation(self):
        record={'archivo':'entrada/problema.txt','contenido':'Newton para sistemas', 'imagenes':[], 'advertencias':[]}
        refs=[{'archivo':'referencias_clase/newton.pdf','contenido':'Jacobiana', 'imagenes':[], 'advertencias':[]},
              {'archivo':'referencias_clase/curso.pdf','contenido':'material extenso','imagenes':['/tmp/imagen_'+str(i)+'.png' for i in range(51)], 'advertencias':[]}]
        with patch('src.codex.invoke',return_value={'archivos':['referencias_clase/newton.pdf'],'motivo':'Newton documentado'}) as ai:
            text,images=main.select_context(ROOT,[record],refs)
        self.assertEqual(ai.call_count,1)
        self.assertIn('Jacobiana',text)
        self.assertEqual(images,[])
        self.assertIn('INVENTARIO',text)


    def test_cli_stdin_schema_read_only_and_invalid_output(self):
        def fake_run(command,**kwargs):
            self.assertIn('read-only',command)
            self.assertIn('--ephemeral',command)
            self.assertEqual(command[-1],'-')
            self.assertEqual(kwargs['input'],'prueba')
            target=Path(command[command.index('-o')+1])
            target.write_text(json.dumps({'reglas':[],'advertencias':[]}))
            return subprocess.CompletedProcess(command,0,'','')
        with patch('src.codex.shutil.which',return_value='/test/codex'), patch('src.codex.subprocess.run',side_effect=fake_run):
            self.assertEqual(codex.invoke(ROOT,'prueba',codex.PROFILE),{'reglas':[],'advertencias':[]})


SHELL=shutil.which('pwsh') or shutil.which('powershell')
@unittest.skipUnless(SHELL,'PowerShell no instalado; contraseña/comandos/limpieza pendientes de validación real')
class PowerShellTests(unittest.TestCase):
    def launch(self,path):
        return subprocess.run([SHELL,'-NoProfile','-File',str(path)],capture_output=True,text=True,timeout=30)

    def test_correct_and_incorrect_password(self):
        from servidor.config import password_record, authorized
        verifier=password_record('prueba-correcta')
        self.assertTrue(authorized('prueba-correcta',verifier))
        self.assertFalse(authorized('incorrecta',verifier))

    def test_four_commands_reject_before_dependencies(self):
        # Las tres operaciones dependen del servidor; limpiar permanece local y no conecta.
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            (root/'cliente').mkdir()
            for name in ('revisar','respuestas','generar','limpiar'):
                shutil.copy2(ROOT/(name+'.ps1'),root/(name+'.ps1'))
                shutil.copy2(ROOT/'cliente'/(name+'.ps1'),root/'cliente'/(name+'.ps1'))
            shutil.copy2(ROOT/'cliente/comun.ps1',root/'cliente/comun.ps1')
            for name in ('revisar','respuestas','generar'):
                result=self.launch(root/(name+'.ps1'))
                self.assertNotEqual(result.returncode,0)
                self.assertNotIn('Consultando Codex',result.stdout)

    def test_cleanup_manifest_and_personal_data(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);fixture(root)
            (root/'sistema').mkdir()
            (root/'src').mkdir()
            own=root/'src/main.py';own.write_text('componente propio')
            altered=root/'README.md';altered.write_text('modificado por usuario')
            academic=root/'entrada/trabajo.m';academic.write_text('x = 1')
            outside=root/'personal.txt';outside.write_text('NO BORRAR')
            manifest={'proyecto':'metodos-cliente-servidor-v2','archivos':[
                {'ruta':'src/main.py','sha256':archivos.digest(own)},
                {'ruta':'README.md','sha256':'0'*64},
                {'ruta':'entrada/trabajo.m','sha256':archivos.digest(academic)},
                {'ruta':'personal.txt','sha256':archivos.digest(outside)}]}
            (root/'sistema/manifiesto.json').write_text(json.dumps(manifest))
            runner=root/'runner.ps1'
            runner.write_text("& '"+str(ROOT/'sistema/desinstalar.ps1').replace("'","''")+"' -Raiz '"+str(root).replace("'","''")+"'",encoding='utf-8')
            result=self.launch(runner)
            self.assertEqual(result.returncode,0,result.stderr)
            self.assertFalse(own.exists())
            self.assertTrue(altered.exists())
            self.assertTrue(academic.exists())
            self.assertTrue(outside.exists())


if __name__=='__main__':
    unittest.main()
